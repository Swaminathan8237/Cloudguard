"""
Script to fill 'Project Work - Data Dictionary.docx' with verified project information.
"""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def get_doc_paths():
    root = Path(__file__).resolve().parent.parent
    if not (root / "docs").exists():
        root = Path.cwd()
        
    backup_candidates = [
        root / "docs" / "Project Work - Data Dictionary_backup.docx",
        Path("docs/Project Work - Data Dictionary_backup.docx"),
        Path("Project Work - Data Dictionary_backup.docx")
    ]
    backup_path = None
    for p in backup_candidates:
        if p.exists():
            backup_path = p
            break
            
    target_path = root / "docs" / "Project Work - Data Dictionary.docx"
    return str(backup_path), str(target_path)

def fill_document():
    backup_path, target_path = get_doc_paths()
    print(f"Loading template from: {backup_path}")
    doc = docx.Document(backup_path)
    
    # 1. Fill student and project details
    # P0: Unique id:
    doc.paragraphs[0].runs[0].text = "Unique id: "
    r = doc.paragraphs[0].add_run("E0224033")
    r.bold = False
    
    # P1: Name:
    doc.paragraphs[1].runs[0].text = "Name: "
    r = doc.paragraphs[1].add_run("Swaminathan")
    r.bold = False
    
    # P3: Project Title:
    doc.paragraphs[3].runs[0].text = "Project Title: "
    r = doc.paragraphs[3].add_run("CloudGuard-Water: Proactive Water Quality Deterioration Prediction Using IoT Data and Machine Learning")
    r.bold = False
    
    # P6: Dataset Description:
    doc.paragraphs[6].runs[0].text = "Dataset Description:"
    doc.paragraphs[6].runs[0].bold = True
    
    desc_text = (
        "This dataset contains 100,800 in-situ physicochemical water quality observations captured using an "
        "IoT-based multi-sensor telemetry system deployed across diverse natural and treated water sources in "
        "Kenya during February–March 2026. The measurements were recorded continuously at a 2-minute sampling "
        "interval across 10 distinct 14-day monitoring campaigns (10,080 readings per deployment). "
        "The system monitored four key physicochemical parameters: pH, Total Dissolved Solids (TDS in ppm), "
        "Turbidity (in NTU), and Temperature (in °C). The dataset covers four contrasting water source categories: "
        "treated municipal and commercial bottled water (Safe), surface water from rivers and lakes subject to silt "
        "and runoff (Unsafe), mineral-rich groundwater from boreholes (Unsafe), and coastal marine water from the "
        "Indian Ocean (Unsafe). Each record includes a binary potability label ('Safe' vs. 'Unsafe') established "
        "against World Health Organization (WHO) and Kenyan national drinking water quality standards."
    )
    doc.paragraphs[7].text = desc_text
    
    # Clear out empty paragraphs P8, P9, P10
    doc.paragraphs[8].text = ""
    doc.paragraphs[9].text = ""
    doc.paragraphs[10].text = ""
    
    # P11: Dataset Download Link:
    doc.paragraphs[11].runs[0].text = "Dataset Download Link: "
    r = doc.paragraphs[11].add_run("https://data.mendeley.com/datasets/575ctg3nnr/2")
    r.bold = False
    
    doc.paragraphs[12].text = "Repository: Mendeley Data (Elsevier) | DOI: 10.17632/575ctg3nnr.2 (Version 2, April 2026)"
    doc.paragraphs[13].text = ""
    doc.paragraphs[14].text = ""
    doc.paragraphs[15].text = ""
    
    # P16: Data Dictionary:
    doc.paragraphs[16].runs[0].text = "Data Dictionary:"
    doc.paragraphs[16].runs[0].bold = True
    
    # P17: Shape of the Data:
    doc.paragraphs[17].text = "Shape of the Data: "
    r = doc.paragraphs[17].add_run("(100800, 5)")
    r.bold = False
    
    # P18: No of Rows
    doc.paragraphs[18].text = "No of Rows: "
    r = doc.paragraphs[18].add_run("100,800")
    r.bold = False
    
    # P19: No of Columns
    doc.paragraphs[19].text = "No of Columns: "
    r = doc.paragraphs[19].add_run("5 (4 features + 1 target label)")
    r.bold = False
    
    # P20: Total Storage Size:
    doc.paragraphs[20].text = "Total Storage Size: "
    r = doc.paragraphs[20].add_run("2.83 MB (2,964,409 bytes on disk; 9.07 MB in RAM)")
    r.bold = False
    
    # 2. Populate the Table
    table = doc.tables[0]
    
    # Table Header
    headers = ["Features (Variable)", "Type Variable", "Description"]
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h
        p = cell.paragraphs[0]
        if p.runs:
            p.runs[0].bold = True
            
    # Data rows
    row_data = [
        (
            "pH",
            "Continuous Numerical (float64)",
            "Logarithmic measure of hydrogen-ion concentration indicating water acidity or alkalinity (0–14 scale). WHO potable drinking water guidelines recommend a target range of 6.5 to 8.5."
        ),
        (
            "TDS_ppm",
            "Continuous Numerical (float64)",
            "Total Dissolved Solids measured in parts per million (ppm = mg/L), quantifying aggregate dissolved inorganic salts and minerals. The desirable drinking water threshold is below 300–500 ppm."
        ),
        (
            "Turbidity_NTU",
            "Continuous Numerical (float64)",
            "Nephelometric Turbidity Units measuring light scattering caused by suspended sediments, silt, and colloidal matter. The standard potable threshold is below 5.0 NTU (ideally < 1.0 NTU)."
        ),
        (
            "Temperature_C",
            "Continuous Numerical (float64)",
            "Ambient water temperature in degrees Celsius (°C). Influences chemical dissolution equilibria, microbial reaction kinetics, and electrical conductivity calibration of sensors."
        ),
        (
            "Label (Target)",
            "Binary Categorical (object / string)",
            "Ground-truth water potability classification ('Safe' vs. 'Unsafe'). 'Safe' represents treated municipal and bottled water (30.0%); 'Unsafe' represents untreated surface, borehole, and marine sources (70.0%)."
        ),
        (
            "pH_change / TDS_change",
            "Continuous Numerical (float64, Engineered)",
            "First-order rate of change of sensor values over the preceding observation window (e.g., TDS_t - TDS_t-k), capturing dynamic deterioration trends and sudden runoff spikes."
        ),
        (
            "rolling_mean / rolling_std",
            "Continuous Numerical (float64, Engineered)",
            "Rolling window statistical features (computed over 5 to 10 observations = 10–20 min), quantifying short-term sensor volatility and progressive contamination momentum."
        ),
        (
            "unsafe_in_next_20_minutes",
            "Binary Categorical (int64, Engineered Target)",
            "Proactive early-warning prediction target indicating whether water will deteriorate into an unsafe condition 20 minutes into the future (enabling automated preventative valve actuation)."
        )
    ]
    
    for i, data in enumerate(row_data):
        row = table.rows[i + 1]
        for j, text in enumerate(data):
            cell = row.cells[j]
            cell.text = text
            if j == 0:
                cell.paragraphs[0].runs[0].bold = True
                
    doc.save(target_path)
    print(f"Successfully saved to: {target_path}")

if __name__ == "__main__":
    fill_document()
